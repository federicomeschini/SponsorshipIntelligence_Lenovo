"""D2b ingestion of the client-supplied FIFA-competition media-value export.

Media equivalency is an AVE-style monetary valuation of exposure. Under P1/W-08
it is banned as a model input: it is staged only for the "industry-claimed
value" comparison line and must never feed an index, attribution or valuation.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.csv as pacsv
import pyarrow.parquet as pq


COLUMNS = {
    "Brand": "brand_line",
    "Date": "date",
    "League/Series": "competition",
    "Market": "market",
    "Region": "region",
    "100% Media Equivalency": "media_value_full",
    "QI Media Value": "media_value_qi",
}


def ingest_media_values(source_path: str, output_dir: str) -> dict:
    """Stage the long export and a competition x month comparison summary."""
    source = Path(source_path)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise ValueError("D2b media-value contract violation: expected one sheet")
    values = list(workbook.worksheets[0].iter_rows(values_only=True))
    workbook.close()
    header = [str(name).strip() for name in values[0]]
    if header != list(COLUMNS):
        raise ValueError(f"D2b media-value contract violation: unexpected columns {header}")

    rows = []
    for record in values[1:]:
        if all(cell is None for cell in record):
            continue
        row = {COLUMNS[name]: cell for name, cell in zip(header, record)}
        day = row["date"]
        row["date"] = day.date() if isinstance(day, datetime) else date.fromisoformat(str(day)[:10])
        for key in ("media_value_full", "media_value_qi"):
            value = float(row[key] or 0.0)
            if value < 0:
                raise ValueError(f"D2b media-value contract violation: negative {key}")
            row[key] = value
        if row["media_value_qi"] > row["media_value_full"] + 1e-9:
            raise ValueError("D2b media-value contract violation: QI value exceeds 100% equivalency")
        row["week"] = row["date"] - timedelta(days=row["date"].weekday())
        row["banned_as_feature"] = True
        rows.append(row)
    rows.sort(key=lambda row: (row["date"], row["competition"], row["market"], row["brand_line"]))

    schema = pa.schema([
        ("brand_line", pa.string()), ("date", pa.date32()), ("competition", pa.string()),
        ("market", pa.string()), ("region", pa.string()), ("media_value_full", pa.float64()),
        ("media_value_qi", pa.float64()), ("week", pa.date32()), ("banned_as_feature", pa.bool_()),
    ])
    table = pa.Table.from_pylist(rows, schema=schema)
    target = Path(output_dir)
    target.mkdir(parents=True, exist_ok=True)
    pq.write_table(table, target / "media_values_daily.parquet")
    pacsv.write_csv(table, target / "media_values_daily.csv")

    summary: dict[tuple[str, str], dict] = {}
    for row in rows:
        key = row["competition"], row["date"].strftime("%Y-%m")
        item = summary.setdefault(key, {"competition": key[0], "month": key[1], "markets": set(),
                                        "media_value_full": 0.0, "media_value_qi": 0.0})
        item["markets"].add(row["market"])
        item["media_value_full"] += row["media_value_full"]
        item["media_value_qi"] += row["media_value_qi"]
    summary_rows = [{**item, "markets": len(item["markets"])}
                    for _, item in sorted(summary.items(), key=lambda pair: (pair[0][1], pair[0][0]))]
    pacsv.write_csv(pa.Table.from_pylist(summary_rows), target / "media_values_competition_monthly.csv")

    total_full = sum(row["media_value_full"] for row in rows)
    total_qi = sum(row["media_value_qi"] for row in rows)
    by_competition = {}
    for item in summary_rows:
        entry = by_competition.setdefault(item["competition"], {"media_value_full": 0.0, "media_value_qi": 0.0})
        entry["media_value_full"] += item["media_value_full"]
        entry["media_value_qi"] += item["media_value_qi"]
    manifest = {
        "source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": "staged_comparison_only", "banned_as_feature": True,
        "permitted_use": "industry-claimed (AVE) comparison line only; never a model input (P1, W-08)",
        "row_count": len(rows),
        "date_min": rows[0]["date"].isoformat(), "date_max": rows[-1]["date"].isoformat(),
        "filename_period_label": "Oct24-Apr26",
        "observed_period_note": "Dated rows start in May 2025; the October 2024 announcement period is not covered.",
        "currency": "not stated in the export",
        "measures": {
            "media_value_full": "100% media equivalency (unadjusted AVE)",
            "media_value_qi": "quality-index adjusted media value",
        },
        "totals": {"media_value_full": round(total_full, 2), "media_value_qi": round(total_qi, 2),
                   "qi_to_full_ratio": round(total_qi / total_full, 4) if total_full else None},
        "by_competition": {name: {key: round(value, 2) for key, value in item.items()}
                           for name, item in sorted(by_competition.items(),
                                                    key=lambda pair: -pair[1]["media_value_qi"])},
        "competitions": sorted({row["competition"] for row in rows}),
        "brand_lines": sorted({row["brand_line"] for row in rows}),
        "markets": len({row["market"] for row in rows}),
    }
    (target / "pull_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", default="data/raw/media_values/lenovo_media_values_fifa_competitions_oct24_apr26.xlsx")
    parser.add_argument("--output-dir", default="data/staged/media_values")
    args = parser.parse_args()
    manifest = ingest_media_values(args.source, args.output_dir)
    print(json.dumps({key: manifest[key] for key in ("row_count", "date_min", "date_max", "totals")}, indent=2))
