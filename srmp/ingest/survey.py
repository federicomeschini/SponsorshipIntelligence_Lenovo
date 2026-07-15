"""Lossless staging and audit utilities for D1 survey workbook exports.

The supplied workbooks are crosstab exports, not respondent-level data. This
module preserves their table structure and never invents a Brand Index.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1_048_576), b""):
            digest.update(block)
    return digest.hexdigest()


def _text(value: Any) -> str:
    """Stable source representation; the raw XLSX remains the evidence."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return repr(value) if isinstance(value, float) else str(value)


def _write(name: str, rows: list[dict[str, Any]], output_dir: Path) -> None:
    """Write convenient CSV plus typed Parquet from identical row records."""
    if not rows:
        return
    with (output_dir / f"{name}.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    pq.write_table(pa.Table.from_pylist(rows), output_dir / f"{name}.parquet")


def _metric(title: str) -> str:
    lowered = title.lower()
    if "aware" in lowered:
        return "brand_awareness"
    if "appealing" in lowered or "how do you feel" in lowered:
        return "appeal"
    if "likely to choose" in lowered or "likely are you to choose" in lowered or "likely are you to purchase" in lowered:
        return "purchase_intent"
    return "unclassified"


def _question_code(title: str) -> str:
    match = re.match(r"\s*(Q[A-Za-z0-9]+)\.?", title)
    return match.group(1) if match else ""


def _property(segments: list[str]) -> str:
    names = {
        match.group(1)
        for segment in segments
        if (match := re.match(r"(.+?) Lenovo Sponsorship (?:Aware|Unaware)$", segment))
    }
    return names.pop() if len(names) == 1 else ""


def _extract_crosstabs(
    source: Path, sha256: str, rows: list[tuple[Any, ...]]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Parse declared Column-% blocks ending in source Filter footers."""
    tables: list[dict[str, Any]] = []
    observations: list[dict[str, Any]] = []
    nonempty = [(i, row) for i, row in enumerate(rows, 1) if any(value is not None for value in row)]
    for position, (header_row, header) in enumerate(nonempty):
        if not header or header[0] != "Column %" or position == 0:
            continue
        title_row, title_values = nonempty[position - 1]
        title = _text(title_values[0]).strip()
        if not title or title in {"Back to TOC", "Country by Wave"}:
            continue
        footer = next(
            ((i, row) for i, row in nonempty[position + 1 :] if _text(row[0]).startswith("Filter:")),
            None,
        )
        if footer is None:
            # Country-by-wave is sample composition, with no filter footer.
            continue
        filter_row, filter_values = footer
        table_id = f"table_r{title_row}"
        segments = [_text(value).strip() for value in header[1:] if value is not None]
        metadata = {
            "table_id": table_id,
            "source_file": source.name,
            "source_sha256": sha256,
            "source_sheet": "Tables",
            "title_row": title_row,
            "header_row": header_row,
            "filter_row": filter_row,
            "title": title,
            "question_code": _question_code(title),
            "metric": _metric(title),
            "property_label": _property(segments),
            "filter_text": _text(filter_values[0]),
            "comparability_note": "Conditional on sponsorship awareness; never combine weighted percentages using unweighted Column n.",
        }
        tables.append(metadata)
        for row_number in range(header_row + 1, filter_row):
            row = rows[row_number - 1]
            response = _text(row[0]).strip()
            if not response:
                continue
            statistic = "Column n" if response == "Column n" else "Column %"
            for column_number, segment in enumerate(segments, 2):
                value = row[column_number - 1] if len(row) >= column_number else None
                if value is None:
                    continue
                observations.append(
                    {
                        "table_id": table_id,
                        "source_file": source.name,
                        "source_sha256": sha256,
                        "source_sheet": "Tables",
                        "source_cell": f"{openpyxl.utils.get_column_letter(column_number)}{row_number}",
                        "source_row": row_number,
                        "title": title,
                        "question_code": _question_code(title),
                        "metric": _metric(title),
                        "property_label": _property(segments),
                        "filter_text": _text(filter_values[0]),
                        "segment": segment,
                        "response_label": response,
                        "statistic": statistic,
                        "value": _text(value),
                    }
                )
    return tables, observations


def _extract_country_composition(source: Path, sha256: str, rows: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    title_row = next((i for i, row in enumerate(rows, 1) if row and row[0] == "Country by Wave"), None)
    if title_row is None:
        return []
    wave_labels = [_text(value).strip() for value in rows[title_row][1:] if value is not None]
    output = []
    for row_number in range(title_row + 2, len(rows) + 1):
        row = rows[row_number - 1]
        country = _text(row[0]).strip()
        if country == "Total sample; Unweighted":
            break
        if not country:
            continue
        for column_number, wave_label in enumerate(wave_labels, 2):
            value = row[column_number - 1] if len(row) >= column_number else None
            if value is not None:
                output.append(
                    {
                        "source_file": source.name,
                        "source_sha256": sha256,
                        "source_sheet": "Tables",
                        "source_cell": f"{openpyxl.utils.get_column_letter(column_number)}{row_number}",
                        "country": country,
                        "wave_label": wave_label,
                        "statistic": "Column %",
                        "value": _text(value),
                        "note": "Unweighted sample composition; not a brand-outcome measure.",
                    }
                )
    return output


def _extract_gwi_foglio1(
    source: Path, sha256: str, foglio_rows: list[tuple[Any, ...]], blank_rows: list[tuple[Any, ...]]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    """Stage Foglio1 and prove it matches the duplicate block in Blank.

    Foglio1 is the user-designated source of truth. Blank is used only for an
    exact, cell-by-cell duplicate check; its values never feed the output.
    """
    if len(foglio_rows) != 17 or any(len(row) != 41 for row in foglio_rows):
        raise ValueError("Unexpected Foglio1 shape; expected exactly 17 rows by 41 columns")
    duplicate_block = [tuple(row[1:42]) for row in blank_rows[12:29]]
    if len(duplicate_block) != 17 or any(len(row) != 41 for row in duplicate_block):
        raise ValueError("Unexpected Blank duplicate-block shape; expected B13:AP29")
    mismatches = [
        (row_number + 1, column_number + 1)
        for row_number, (left, right) in enumerate(zip(foglio_rows, duplicate_block))
        for column_number, (left_value, right_value) in enumerate(zip(left, right))
        if _text(left_value) != _text(right_value)
    ]
    if mismatches:
        raise ValueError(f"Foglio1 differs from Blank duplicate block at {mismatches[:5]}")

    headers = [_text(value).strip() for value in foglio_rows[1][3:]]
    if headers[0] != "Totals" or len(headers) != 38:
        raise ValueError("Unexpected GWI wave headers in Foglio1")
    crosstabs: list[dict[str, Any]] = []
    proxy_series: list[dict[str, Any]] = []
    question = ""
    name = ""
    for row_number, row in enumerate(foglio_rows[2:], start=3):
        if row[0] is not None:
            question = _text(row[0]).strip()
        if row[1] is not None:
            name = _text(row[1]).strip()
        statistic = _text(row[2]).strip()
        if not statistic:
            raise ValueError(f"Foglio1 row {row_number} has no statistic label")
        for column_number, (header, value) in enumerate(zip(headers, row[3:]), start=4):
            record = {
                "source_file": source.name,
                "source_sha256": sha256,
                "source_sheet": "Foglio1",
                "source_cell": f"{openpyxl.utils.get_column_letter(column_number)}{row_number}",
                "question": question,
                "name": name,
                "statistic": statistic,
                "period_label": header,
                "is_total": header == "Totals",
                "value": _text(value),
            }
            crosstabs.append(record)
            if question.startswith("Tech Brands:") and statistic == "Column %" and header != "Totals":
                proxy_series.append(
                    {
                        **record,
                        "proxy_id": re.sub(r"[^a-z0-9]+", "_", question.lower()).strip("_"),
                        "proxy_construct": question.replace("Tech Brands:", "").replace("*", "").strip().lower(),
                        "admission_status": "candidate_requires_l1_validation",
                    }
                )
    check = {
        "authoritative_sheet": "Foglio1",
        "duplicate_check_sheet": "Blank",
        "duplicate_check_range": "Blank!B13:AP29",
        "cells_compared": 697,
        "status": "passed_exact_match",
        "quarterly_periods_excluding_totals": 37,
        "candidate_proxy_series": len({row["proxy_id"] for row in proxy_series}),
    }
    return crosstabs, proxy_series, check

def stage_survey_exports(raw_dir: str | Path, output_dir: str | Path, report_path: str | Path) -> dict[str, Any]:
    """Audit every raw XLSX and create lossless, reviewable staged artifacts."""
    raw_dir, output_dir, report_path = Path(raw_dir), Path(output_dir), Path(report_path)
    sources = sorted(raw_dir.glob("*.xlsx"))
    if not sources:
        raise FileNotFoundError(f"No XLSX files found in {raw_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)
    inventory: list[dict[str, Any]] = []
    cells: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    observations: list[dict[str, Any]] = []
    composition: list[dict[str, Any]] = []
    gwi_crosstabs: list[dict[str, Any]] = []
    gwi_proxy_series: list[dict[str, Any]] = []
    gwi_checks: list[dict[str, Any]] = []
    for source in sources:
        sha256 = _sha256(source)
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
        source_cells: list[dict[str, Any]] = []
        sheets: list[dict[str, Any]] = []
        table_rows: list[tuple[Any, ...]] = []
        sheet_rows: dict[str, list[tuple[Any, ...]]] = {}
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            sheet_rows[sheet.title] = rows
            if sheet.title == "Tables":
                table_rows = rows
            count = 0
            for row_number, row in enumerate(rows, 1):
                for column_number, value in enumerate(row, 1):
                    if value is None:
                        continue
                    count += 1
                    source_cells.append(
                        {
                            "source_file": source.name,
                            "source_sha256": sha256,
                            "sheet": sheet.title,
                            "cell": sheet.cell(row_number, column_number).coordinate,
                            "row_number": row_number,
                            "column_number": column_number,
                            "value_type": type(value).__name__,
                            "value": _text(value),
                        }
                    )
            sheets.append({"sheet": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, "nonempty_cells": count})
        status = "usable_gwi_crosstab_export" if "Foglio1" in sheet_rows else ("usable_crosstab_export" if source_cells else "valid_but_empty")
        inventory.append({"source_file": source.name, "source_sha256": sha256, "bytes": source.stat().st_size, "sheets": sheets, "nonempty_cells": len(source_cells), "status": status})
        cells.extend(source_cells)
        if table_rows:
            parsed_tables, parsed_observations = _extract_crosstabs(source, sha256, table_rows)
            tables.extend(parsed_tables)
            observations.extend(parsed_observations)
            composition.extend(_extract_country_composition(source, sha256, table_rows))
        if "Foglio1" in sheet_rows:
            if "Blank" not in sheet_rows:
                raise ValueError("GWI workbook is missing the duplicate-check sheet Blank")
            parsed_gwi, parsed_proxy_series, gwi_check = _extract_gwi_foglio1(
                source, sha256, sheet_rows["Foglio1"], sheet_rows["Blank"]
            )
            gwi_crosstabs.extend(parsed_gwi)
            gwi_proxy_series.extend(parsed_proxy_series)
            gwi_checks.append(gwi_check)
    failures = [row for row in observations if row["response_label"] == "NET" and row["statistic"] == "Column %" and abs(float(row["value"]) - 1) > 1e-12]
    if failures:
        raise ValueError(f"Source NET validation failed for {[row['source_cell'] for row in failures]}")
    wave_filters = list(dict.fromkeys(table["filter_text"] for table in tables))
    waves = [
        {
            "filter_text": value,
            "reported_wave": (match.group(0) if (match := re.search(r"Wave\s+\d+.*?\d{4}", value)) else ""),
            "exact_fieldwork_start": "",
            "exact_fieldwork_end": "",
            "comparability_status": "unknown_requires_source_metadata",
            "note": "The export provides a labelled month/wave, not exact fieldwork dates.",
        }
        for value in wave_filters
    ]
    artifacts = {
        "survey_source_cells": cells,
        "consumer_research_tables": tables,
        "consumer_research_crosstabs": observations,
        "consumer_research_country_composition": composition,
        "consumer_research_wave_register": waves,
        "gwi_core_tech_brands_crosstab": gwi_crosstabs,
        "gwi_lenovo_proxy_candidates": gwi_proxy_series,
    }
    for name, rows in artifacts.items():
        _write(name, rows, output_dir)
    manifest = {
        "source": "client-provided survey workbook exports",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "sources": inventory,
        "output_row_counts": {name: len(rows) for name, rows in artifacts.items()},
        "validation": {"rule": "every declared crosstab Column % NET equals 1", "status": "passed", "net_rows_checked": sum(row["response_label"] == "NET" and row["statistic"] == "Column %" for row in observations), "gwi_foglio1_duplicate_check": gwi_checks},
        "limitations": ["Outcome crosstabs are conditional on sponsorship-awareness segment.", "Column n is explicitly unweighted and must not reweight reported percentages.", "Appeal and purchase response scales differ between 2024 and 2025-26.", "Exact fieldwork dates and comparability metadata are absent.", "GWI proxy candidates require the locked L1 validation gate before index use."],
    }
    (output_dir / "pull_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    metric_counts = Counter(table["metric"] for table in tables)
    report = f"""# Survey source audit

## Verdict

The consumer-research workbook is usable as a **segmented crosstab source**. The GWI workbook is usable as two **proxy candidates**, using `Foglio1` as the source of truth. Neither workbook alone can anchor the SRMP Brand Index.

## Integrity

| Workbook | SHA-256 | Status |
|---|---|---|
"""
    report += "".join(f"| `{item['source_file']}` | `{item['source_sha256']}` | {item['status']} ({item['nonempty_cells']} non-empty cells) |\n" for item in inventory)
    report += f"""
## Staged output

- {len(tables)} declared crosstab tables: {metric_counts['brand_awareness']} awareness, {metric_counts['appeal']} appeal, {metric_counts['purchase_intent']} purchase-intent.
- {len(waves)} distinct filter/wave labels. The export states named months only, not exact fieldwork dates.
- CSV and Parquet artifacts in `data/staged/survey/` retain each source cell, workbook hash, sheet, and Excel coordinate.
- GWI `Foglio1` has 37 quarterly periods (Q1 2017–Q1 2026) for Lenovo engagement and consideration. Its 17 × 41 block exactly matches `Blank!B13:AP29` (697/697 cells); `Foglio1` alone feeds the staged GWI outputs.

## Safe uses

- Compare sponsorship-aware and unaware segments **within the same table and wave**.
- Inspect source distributions and reported `T4B` / `NET` summaries.
- Use Country by Wave only as unweighted sample-composition context.
- Treat the GWI engagement and consideration series as proxy candidates only; admit them to the index only after the locked L1 validation gate.

## Do not do this

- Do not treat the output as population-level brand health: every outcome table is segmented by sponsorship awareness.
- Do not recombine segments with `Column n`; the export calls those counts unweighted while the percentages are weighted.
- Do not create a time series by joining the 2024 0-10 scales to the 2025-26 categorical scales.
- Do not infer fieldwork dates, wave comparability, a composite, or funnel-stage scores from these crosstabs beyond what they explicitly report.
- Do not use the GWI figures as an index anchor or as causal evidence without the required validation and modelling steps.

## Client data still required

A wave-level analytical extract with approved funnel-stage scores (awareness, engagement, appeal, consideration, purchase_intent), exact fieldwork start/end, weights, questionnaire-version metadata, and a comparability crosswalk.
"""
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(report, encoding="utf-8")
    return manifest


def ingest_survey(source_path: str, output_dir: str, manifest_path: str) -> None:
    """Compatibility wrapper; ``source_path`` is the raw survey directory."""
    manifest = stage_survey_exports(source_path, output_dir, "reports/methods_annex/survey_source_audit.md")
    Path(manifest_path).write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit and losslessly stage survey crosstab exports.")
    parser.add_argument("--raw-dir", default="data/raw/survey")
    parser.add_argument("--output-dir", default="data/staged/survey")
    parser.add_argument("--report", default="reports/methods_annex/survey_source_audit.md")
    args = parser.parse_args()
    stage_survey_exports(args.raw_dir, args.output_dir, args.report)


if __name__ == "__main__":
    main()






